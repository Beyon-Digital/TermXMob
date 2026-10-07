"""Versioned workspace artifacts and explicit-account media adapters."""
from __future__ import annotations

import base64
import csv
import hashlib
import io
import json
import math
import os
import sqlite3
import uuid
from contextlib import contextmanager
from html import escape
from pathlib import Path
from time import time

import httpx
from fastapi import HTTPException
from termx.audit import log_event

KINDS = {'document', 'table', 'chart', 'deck', 'code', 'image', 'audio', 'video', 'input'}
LIMIT = 100 * 1024 * 1024


class ArtifactService:
    def __init__(self, directory):
        self.root = Path(directory)
        self.root.mkdir(parents=True, exist_ok=True, mode=0o700)
        self.path = self.root / 'artifacts.sqlite3'
        with self.db() as db:
            db.executescript('''CREATE TABLE IF NOT EXISTS artifacts (
                id TEXT PRIMARY KEY, owner TEXT, project TEXT, kind TEXT, title TEXT, version INTEGER);
                CREATE TABLE IF NOT EXISTS versions (
                artifact TEXT, version INTEGER, content TEXT, blob TEXT, mime TEXT, created REAL,
                PRIMARY KEY(artifact,version));
                CREATE TABLE IF NOT EXISTS operations (
                owner TEXT, id TEXT, digest TEXT, status TEXT, result TEXT,
                project TEXT, created REAL, PRIMARY KEY(owner,id));''')
            columns = {row['name'] for row in db.execute('PRAGMA table_info(operations)')}
            if 'project' not in columns:
                db.execute('ALTER TABLE operations ADD COLUMN project TEXT')
            if 'created' not in columns:
                db.execute('ALTER TABLE operations ADD COLUMN created REAL')
            # The host cannot establish whether an interrupted provider request
            # was charged. Retain it for inspection, never replay it on restart.
            db.execute("UPDATE operations SET status='unknown' WHERE status='executing'")
        self.path.chmod(0o600)

    @contextmanager
    def db(self):
        connection = sqlite3.connect(self.path, timeout=15)
        connection.row_factory = sqlite3.Row
        try:
            with connection:
                yield connection
        finally:
            connection.close()

    def create(self, owner, project, kind, title, content=None, blob=None, mime=None):
        if kind not in KINDS:
            raise HTTPException(400, 'Unsupported artifact kind')
        if blob is None:
            self.validate(kind, content)
        elif len(blob) > LIMIT:
            raise HTTPException(413, 'Media exceeds 100 MiB')
        identifier = uuid.uuid4().hex
        with self.db() as db:
            db.execute('INSERT INTO artifacts VALUES(?,?,?,?,?,0)', (identifier, owner, project, kind, title[:200]))
        return self.revise(owner, identifier, 0, content=content, blob=blob, mime=mime)

    def get(self, owner, identifier, version=None):
        with self.db() as db:
            artifact = db.execute('SELECT * FROM artifacts WHERE id=? AND owner=?', (identifier, owner)).fetchone()
            if not artifact:
                raise HTTPException(404, 'Artifact not found')
            item = db.execute('SELECT * FROM versions WHERE artifact=? AND version=?', (identifier, version or artifact['version'])).fetchone()
            if not item:
                raise HTTPException(404, 'Artifact version not found')
            result = {**dict(artifact), **dict(item)}
            result['content'] = json.loads(result['content']) if result['content'] else None
            result.pop('blob')
            return result

    def list(self, owner, project=None):
        with self.db() as db:
            rows = db.execute('SELECT * FROM artifacts WHERE owner=? ORDER BY rowid DESC LIMIT 500', (owner,)).fetchall()
        return [dict(row) for row in rows if project is None or row['project'] == project]

    def operations(self, owner, project):
        with self.db() as db:
            rows = db.execute('SELECT id,status,result,created FROM operations WHERE owner=? AND project IS ? ORDER BY rowid DESC LIMIT 100', (owner, project)).fetchall()
        return [{**dict(row), 'result': {key: value for key,value in json.loads(row['result']).items() if key in {'id','title','kind','version'}} if row['result'] else None} for row in rows]

    def validate(self, kind, content):
        encoded = json.dumps(content)
        if len(encoded.encode()) > 2_000_000:
            raise HTTPException(413, 'Artifact content exceeds two megabytes')
        if kind in {'document', 'code'} and not isinstance(content, str):
            raise HTTPException(400, 'Text artifact content must be a string')
        if kind == 'table':
            if not isinstance(content, dict) or set(content) - {'columns', 'rows'} or not isinstance(content.get('columns'), list) or not isinstance(content.get('rows'), list):
                raise HTTPException(400, 'Table requires columns and rows')
            if len(content['columns']) > 200 or len(content['rows']) > 10000:
                raise HTTPException(413, 'Table exceeds 200 columns or 10000 rows')
            if any(not isinstance(row, list) or len(row) != len(content['columns']) for row in content['rows']):
                raise HTTPException(400, 'Table rows must match columns')
        if kind == 'chart':
            if not isinstance(content, dict) or not isinstance(content.get('points'), list) or len(content['points']) > 1000:
                raise HTTPException(400, 'Chart requires up to 1000 labelled points')
            for point in content['points']:
                if not isinstance(point, dict) or not isinstance(point.get('value'), (int, float)) or not math.isfinite(point['value']) or not isinstance(point.get('label'), str):
                    raise HTTPException(400, 'Chart points require a label and numeric value')
        if kind == 'deck':
            if not isinstance(content, list) or len(content) > 100 or any(not isinstance(s, dict) or set(s) - {'title', 'body', 'notes'} for s in content):
                raise HTTPException(400, 'Deck requires up to 100 title/body/notes slides')

    def revise(self, owner, identifier, expected, *, content=None, blob=None, mime=None):
        with self.db() as db:
            db.execute('BEGIN IMMEDIATE')
            artifact = db.execute('SELECT * FROM artifacts WHERE id=? AND owner=?', (identifier, owner)).fetchone()
            if not artifact:
                raise HTTPException(404, 'Artifact not found')
            if expected != artifact['version']:
                raise HTTPException(409, 'Artifact changed in another window; reload before saving')
            if blob is None:
                self.validate(artifact['kind'], content)
                target = None
            else:
                if len(blob) > LIMIT:
                    raise HTTPException(413, 'Media exceeds 100 MiB')
                target = self.root / (hashlib.sha256(blob).hexdigest() + '.blob')
                if not target.exists():
                    fd = os.open(target, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o600)
                    with os.fdopen(fd, 'wb') as handle:
                        handle.write(blob)
            version = expected + 1
            db.execute('INSERT INTO versions VALUES(?,?,?,?,?,?)', (identifier, version, json.dumps(content) if content is not None else None, str(target) if target else None, mime, time()))
            db.execute('UPDATE artifacts SET version=? WHERE id=?', (version, identifier))
        log_event('artifact_revision', artifact_id=identifier, version=version, artifact_kind=artifact['kind'])
        return self.get(owner, identifier)

    def versions(self, owner, identifier):
        self.get(owner, identifier)
        with self.db() as db:
            return [dict(row) for row in db.execute('SELECT version, mime, created FROM versions WHERE artifact=? ORDER BY version DESC', (identifier,))]

    def binary(self, owner, identifier, version=None):
        item = self.get(owner, identifier, version)
        with self.db() as db:
            row = db.execute('SELECT blob FROM versions WHERE artifact=? AND version=?', (identifier, item['version'])).fetchone()
        if not row['blob']:
            raise HTTPException(400, 'Artifact is not binary')
        return Path(row['blob']).read_bytes(), item['mime'] or 'application/octet-stream'

    def export(self, owner, identifier, format, version=None):
        item = self.get(owner, identifier, version)
        kind, value = item['kind'], item['content']
        stream = io.BytesIO()
        if format == 'json':
            return json.dumps(value, ensure_ascii=False, indent=2).encode(), 'application/json'
        if format in {'txt', 'md', 'code'} and kind in {'document', 'code'}:
            return value.encode(), 'text/plain; charset=utf-8'
        if format == 'docx' and kind == 'document':
            from docx import Document
            document = Document()
            for paragraph in value.split('\n'):
                document.add_paragraph(paragraph)
            document.save(stream)
            return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
        if format in {'csv', 'xlsx'} and kind == 'table':
            rows = [value['columns'], *value['rows']]
            if format == 'csv':
                text = io.StringIO(newline='')
                # Spreadsheet formulas supplied as text remain text on export.
                csv.writer(text).writerows([["'" + v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v for v in row] for row in rows])
                return text.getvalue().encode(), 'text/csv; charset=utf-8'
            from openpyxl import Workbook
            book = Workbook()
            for row in rows:
                book.active.append(row)
            for row in book.active:
                for cell in row:
                    if isinstance(cell.value,str):
                        cell.data_type = 's'
            book.save(stream)
            return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        if format == 'pptx' and kind == 'deck':
            from pptx import Presentation
            deck = Presentation()
            for content in value:
                slide = deck.slides.add_slide(deck.slide_layouts[1])
                slide.shapes.title.text = str(content.get('title',''))
                slide.placeholders[1].text = str(content.get('body',''))
                slide.notes_slide.notes_text_frame.text = str(content.get('notes',''))
            deck.save(stream)
            return stream.getvalue(), 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
        if format == 'svg' and kind == 'chart':
            points = value['points']
            scale = max([abs(p['value']) for p in points] or [1]) or 1
            height = max(100, len(points)*32+56)
            bars = ''.join(f'<text x="12" y="{i*32+46}">{escape(p["label"])}</text><rect x="160" y="{i*32+28}" width="{abs(p["value"])/scale*400:.2f}" height="20" fill="#116c46"/><text x="570" y="{i*32+46}">{p["value"]}</text>' for i,p in enumerate(points))
            svg = f'<svg xmlns="http://www.w3.org/2000/svg" width="720" height="{height}" role="img"><title>{escape(item["title"])}</title><rect width="720" height="{height}" fill="white"/><g fill="#191f1c" font-family="sans-serif" font-size="14">{bars}</g></svg>'
            return svg.encode(), 'image/svg+xml'
        if format == 'original' and kind in {'input','image','audio','video'}:
            return self.binary(owner, identifier, version)
        raise HTTPException(400, 'This export format is unavailable for the artifact kind')


class MediaProvider:
    """OpenAI-compatible image/audio API, explicitly selected existing account.

    No subscription/API-key fallback. The caller must disclose and acknowledge
    the selected provider/model billing before making a generation request.
    """
    def __init__(self, state, artifacts, client=None):
        self.state, self.artifacts, self.client = state, artifacts, client

    @staticmethod
    async def bounded_post(client, url, **kwargs):
        # Streaming enforces the bound before allocating an arbitrary provider
        # response. Cancellation closes the connection without retrying billing.
        async with client.stream('POST', url, **kwargs) as response:
            if response.status_code >= 300:
                raise HTTPException(409, 'Provider rejected this operation; verify model capability, entitlement and quota')
            output = bytearray()
            async for chunk in response.aiter_bytes():
                if len(output) + len(chunk) > LIMIT:
                    raise HTTPException(413, 'Provider output exceeds the media limit')
                output.extend(chunk)
            return httpx.Response(response.status_code, content=bytes(output), headers=response.headers)

    async def generate(self, owner, project, request):
        if not request.get('acknowledge_billing'):
            raise HTTPException(409, 'Confirm this provider/model usage before generation')
        provider = self.state.agent_store.get_provider(request['provider_id'])
        if not provider or provider['kind'] != 'openai-compatible':
            raise HTTPException(400, 'Select a configured compatible API account')
        operation = request['operation']
        capability = 'image' if operation.startswith('image') else 'audio'
        if capability not in provider.get('capabilities', []):
            raise HTTPException(409, f'The selected account does not declare {capability} support')
        key = self.state.credentials.get(provider['id'])
        if not key:
            raise HTTPException(401, 'Connect this provider account on the host')
        model = request.get('model')
        if not isinstance(model, str) or not model.strip():
            raise HTTPException(400, 'Select an explicit supported model')
        digest = hashlib.sha256(json.dumps({'project':project,'request':request},sort_keys=True).encode()).hexdigest()
        request_id = request['request_id']
        with self.artifacts.db() as db:
            db.execute('BEGIN IMMEDIATE')
            prior = db.execute('SELECT * FROM operations WHERE owner=? AND id=?', (owner,request_id)).fetchone()
            if prior:
                if prior['digest'] != digest:
                    raise HTTPException(409, 'Media action ID reused with different arguments')
                if prior['result']:
                    return json.loads(prior['result'])
                raise HTTPException(409, 'Previous provider outcome requires reconciliation; billing will not be replayed')
            db.execute("INSERT INTO operations(owner,id,digest,status,result,project,created) VALUES(?,?,?,'executing',NULL,?,?)", (owner,request_id,digest,project,time()))
        client = self.client or httpx.AsyncClient(timeout=120, follow_redirects=False)
        base = provider['base_url'].rstrip('/').removesuffix('/responses')
        headers = {'Authorization': f'Bearer {key}'}
        try:
            if operation == 'image-generate':
                payload = {'model':model, 'prompt':request['prompt'], 'n':1}
                if model.startswith('dall-e'):
                    payload['response_format'] = 'b64_json'
                response = await self.bounded_post(client,base+'/images/generations', headers=headers, json=payload)
            elif operation in {'image-edit','transcribe'}:
                source = self.artifacts.get(owner,request['input_id'])
                if source['project'] != project:
                    raise HTTPException(403, 'Input belongs to another project')
                raw, mime = self.artifacts.binary(owner,request['input_id'])
                if operation == 'image-edit' and not mime.startswith('image/'):
                    raise HTTPException(400, 'Image editing requires an image')
                if operation == 'transcribe' and not mime.startswith('audio/'):
                    raise HTTPException(400, 'Transcription requires audio')
                response = await self.bounded_post(client,base+('/images/edits' if operation == 'image-edit' else '/audio/transcriptions'), headers=headers,
                    data={'model':model, **({'prompt':request['prompt']} if operation == 'image-edit' else {})},
                    files={'image' if operation=='image-edit' else 'file':('input.'+{'audio/webm':'webm','audio/ogg':'ogg','audio/mp4':'m4a','audio/mpeg':'mp3','image/jpeg':'jpg','image/webp':'webp'}.get(mime,'png' if operation=='image-edit' else 'wav'), raw,mime)})
            elif operation == 'speak':
                response = await self.bounded_post(client,base+'/audio/speech',headers=headers,json={'model':model,'input':request['prompt'],'voice':request.get('voice','alloy'),'response_format':'mp3'})
            else:
                raise HTTPException(400,'Unsupported media operation')
            if response.status_code >= 300:
                raise HTTPException(409,'Provider rejected this operation; verify model capability, entitlement and quota')
            if len(response.content)>LIMIT:
                raise HTTPException(413,'Provider output exceeds the media limit')
            if operation == 'transcribe':
                result = self.artifacts.create(owner,project,'document','Transcript',response.json()['text'])
            elif operation == 'speak':
                result = self.artifacts.create(owner,project,'audio','Spoken response',blob=response.content,mime='audio/mpeg')
            else:
                raw = base64.b64decode(response.json()['data'][0]['b64_json'],validate=True)
                from PIL import Image
                image = Image.open(io.BytesIO(raw))
                image.verify()
                result = self.artifacts.create(owner,project,'image','Generated image',blob=raw,mime=Image.MIME.get(image.format,'image/png'))
            with self.artifacts.db() as db:
                db.execute('UPDATE operations SET status=\'completed\',result=? WHERE owner=? AND id=?',(json.dumps(result),owner,request_id))
            return result
        except httpx.RequestError as exc:
            raise HTTPException(503,'Provider unavailable; inspect the operation before retrying') from exc
        except (KeyError, IndexError, ValueError, TypeError) as exc:
            raise HTTPException(502,'Provider output did not match the selected capability; the operation will not be replayed') from exc
        finally:
            with self.artifacts.db() as db:
                db.execute("UPDATE operations SET status='unknown' WHERE owner=? AND id=? AND status='executing'", (owner,request_id))
            if self.client is None:
                await client.aclose()
