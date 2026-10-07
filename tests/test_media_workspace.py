import asyncio
import base64
import io
from types import SimpleNamespace
from zipfile import ZipFile
import httpx
import pytest
from fastapi import HTTPException
from termx.media.service import ArtifactService, MediaProvider
from termx.media.http import input_context


def test_versioned_documents_tables_decks_chart_exports(tmp_path):
    service=ArtifactService(tmp_path)
    document=service.create('alice','project','document','Notes','First\nSecond')
    data,mime=service.export('alice',document['id'],'docx')
    assert 'word/document.xml' in ZipFile(io.BytesIO(data)).namelist()
    newer=service.revise('alice',document['id'],1,content='Revised')
    assert newer['version']==2
    assert service.get('alice',document['id'],1)['content']=='First\nSecond'
    with pytest.raises(HTTPException):service.revise('alice',document['id'],1,content='Conflict')
    with pytest.raises(HTTPException):service.get('bob',document['id'])
    table=service.create('alice','project','table','Data',{'columns':['Name','Total'],'rows':[['=danger',4],['B',5]]})
    data,_=service.export('alice',table['id'],'xlsx')
    from openpyxl import load_workbook
    book=load_workbook(io.BytesIO(data))
    assert book.active['A2'].value=='=danger' and book.active['A2'].data_type=='s'
    csv,_=service.export('alice',table['id'],'csv')
    assert "'=danger" in csv.decode()
    deck=service.create('alice','project','deck','Slides',[{'title':'Title','body':'Content','notes':'Speaker note'}])
    data,_=service.export('alice',deck['id'],'pptx')
    assert 'ppt/slides/slide1.xml' in ZipFile(io.BytesIO(data)).namelist()
    chart=service.create('alice','project','chart','Chart<script>',{'points':[{'label':'<script>alert(1)</script>','value':42}]})
    svg,_=service.export('alice',chart['id'],'svg')
    assert b'<script>' not in svg and b'&lt;script&gt;' in svg


def test_image_conversion_has_bounded_size_no_metadata(tmp_path):
    from PIL import Image
    stream=io.BytesIO();Image.new('RGB',(2000,100),'red').save(stream,format='PNG')
    service=ArtifactService(tmp_path)
    item=service.create('alice','project','input','Image',blob=stream.getvalue(),mime='image/png')
    result=input_context(service,'alice',item['id'])
    raw=base64.b64decode(result['attachments'][0]['data'])
    image=Image.open(io.BytesIO(raw))
    assert image.width==1600 and len(raw)<2*1024*1024
    text=service.create('alice','project','input','Text',blob=b'Actual text',mime='text/plain')
    assert input_context(service,'alice',text['id'])['text']=='Actual text'


def test_image_generation_explicit_provider_protocol_and_dedup(tmp_path):
    from PIL import Image
    image=io.BytesIO();Image.new('RGB',(10,10),'blue').save(image,format='PNG')
    calls=[]
    def handler(request):
        calls.append(request)
        assert request.url.path=='/v1/images/generations'
        assert request.headers['authorization']=='Bearer test-only'
        return httpx.Response(200,json={'data':[{'b64_json':base64.b64encode(image.getvalue()).decode()}]})
    provider={'id':'configured','kind':'openai-compatible','base_url':'https://provider.example/v1','capabilities':['image','audio']}
    state=SimpleNamespace(agent_store=SimpleNamespace(get_provider=lambda _:provider),credentials=SimpleNamespace(get=lambda _:'test-only'))
    service=ArtifactService(tmp_path)
    client=httpx.AsyncClient(transport=httpx.MockTransport(handler))
    adapter=MediaProvider(state,service,client)
    args={'request_id':'request-123','provider_id':'configured','operation':'image-generate','model':'dall-e-fixture','prompt':'Blue square','acknowledge_billing':True}
    async def run():
        first=await adapter.generate('alice','project',args)
        assert await adapter.generate('alice','project',args)==first
        assert first['kind']=='image' and first['version']==1
        with pytest.raises(HTTPException):await adapter.generate('alice','other-project',args)
        with pytest.raises(HTTPException):await adapter.generate('alice','project',{**args,'request_id':'request-456','acknowledge_billing':False})
        await client.aclose()
    asyncio.run(run())
    assert len(calls)==1


def test_provider_transcription_speech_entitlement_and_no_replay(tmp_path):
    calls=[]
    def handler(request):
        calls.append(request.url.path)
        if request.url.path.endswith('/transcriptions'):return httpx.Response(200,json={'text':'Converted speech'})
        if request.url.path.endswith('/speech'):return httpx.Response(200,content=b'audio-response')
        return httpx.Response(403,json={'error':'missing entitlement'})
    provider={'id':'configured','kind':'openai-compatible','base_url':'https://provider.example/v1','capabilities':['audio','image']}
    state=SimpleNamespace(agent_store=SimpleNamespace(get_provider=lambda _:provider),credentials=SimpleNamespace(get=lambda _:'test-only'))
    service=ArtifactService(tmp_path)
    source=service.create('alice','project','input','Voice',blob=b'fixture-audio',mime='audio/wav')
    client=httpx.AsyncClient(transport=httpx.MockTransport(handler));adapter=MediaProvider(state,service,client)
    args={'request_id':'transcribe-123','provider_id':'configured','operation':'transcribe','model':'audio-fixture','input_id':source['id'],'acknowledge_billing':True}
    async def run():
        transcript=await adapter.generate('alice','project',args)
        assert transcript['content']=='Converted speech'
        audio=await adapter.generate('alice','project',{**args,'request_id':'speech-123','operation':'speak','prompt':'Hello'})
        assert service.binary('alice',audio['id'])[0]==b'audio-response'
        failed={**args,'request_id':'failed-123','operation':'image-generate','prompt':'Test'}
        with pytest.raises(HTTPException,match='Provider rejected'):await adapter.generate('alice','project',failed)
        with pytest.raises(HTTPException,match='will not be replayed'):await adapter.generate('alice','project',failed)
        await client.aclose()
    asyncio.run(run())
    assert len(calls)==3
    unknown=service.operations('alice','project')
    assert unknown[0]['status']=='unknown'
    assert service.operations('bob','project')==[]
    assert service.operations('alice','other-project')==[]
    assert all(not row['result'] or 'content' not in row['result'] for row in unknown)


def test_media_restart_preserves_unknown_outcome_without_replaying(tmp_path):
    import sqlite3
    # Exercise migration of the original five-column durable store too.
    tmp_path.mkdir(exist_ok=True)
    db=sqlite3.connect(tmp_path/'artifacts.sqlite3')
    db.execute('CREATE TABLE operations(owner TEXT,id TEXT,digest TEXT,status TEXT,result TEXT,PRIMARY KEY(owner,id))')
    db.execute("INSERT INTO operations VALUES('alice','old-operation','hash','executing',NULL)")
    db.commit();db.close()
    service=ArtifactService(tmp_path)
    assert service.operations('alice',None)[0]['status']=='unknown'


def test_provider_stream_limit_closes_response_and_preserves_unknown_billing(tmp_path, monkeypatch):
    import termx.media.service as module
    monkeypatch.setattr(module, 'LIMIT', 12)
    class Oversized(httpx.AsyncByteStream):
        closed=False
        async def __aiter__(self):
            yield b'a' * 8
            yield b'b' * 8
            raise AssertionError('Response must stop at the bounded chunk')
        async def aclose(self):self.closed=True
    stream=Oversized();calls=[]
    def handler(request):
        calls.append(request)
        return httpx.Response(200, stream=stream)
    provider={'id':'account','kind':'openai-compatible','base_url':'https://provider.example/v1','capabilities':['audio']}
    state=SimpleNamespace(agent_store=SimpleNamespace(get_provider=lambda _:provider),credentials=SimpleNamespace(get=lambda _:'fixture'))
    artifacts=ArtifactService(tmp_path);client=httpx.AsyncClient(transport=httpx.MockTransport(handler));adapter=MediaProvider(state,artifacts,client)
    args={'request_id':'stream-limit','provider_id':'account','operation':'speak','model':'explicit','prompt':'Text','acknowledge_billing':True}
    async def run():
        with pytest.raises(HTTPException,match='media limit'):await adapter.generate('alice','project',args)
        assert stream.closed
        assert artifacts.operations('alice','project')[0]['status']=='unknown'
        with pytest.raises(HTTPException,match='will not be replayed'):await adapter.generate('alice','project',args)
        await client.aclose()
    asyncio.run(run());assert len(calls)==1


def test_real_pdf_document_video_context_and_expanded_document_limit(tmp_path):
    import subprocess
    import imageio_ffmpeg
    from docx import Document
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    service=ArtifactService(tmp_path/'artifacts')
    writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    contents=DecodedStreamObject();contents.set_data(b'BT /F1 12 Tf 20 200 Td (Real PDF context) Tj ET');page[NameObject('/Contents')]=writer._add_object(contents)
    blob=io.BytesIO();writer.write(blob)
    pdf=service.create('alice','project','input','PDF',blob=blob.getvalue(),mime='application/pdf')
    assert 'Real PDF context' in input_context(service,'alice',pdf['id'])['text']
    doc=Document();doc.add_paragraph('Real document context');blob=io.BytesIO();doc.save(blob)
    mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    document=service.create('alice','project','input','DOCX',blob=blob.getvalue(),mime=mime)
    assert input_context(service,'alice',document['id'])['text']=='Real document context'
    from zipfile import ZIP_DEFLATED
    oversized=io.BytesIO()
    with ZipFile(oversized,'w',compression=ZIP_DEFLATED) as archive:archive.writestr('word/document.xml',b' '* (50*1024*1024+1))
    bad=service.create('alice','project','input','Expanded DOCX',blob=oversized.getvalue(),mime=mime)
    with pytest.raises(HTTPException,match='expanded contents'):input_context(service,'alice',bad['id'])
    video=tmp_path/'video.mp4'
    subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(),'-nostdin','-v','error','-f','lavfi','-i','color=c=blue:s=64x64:r=1:d=6','-c:v','libx264','-pix_fmt','yuv420p',str(video)],capture_output=True,check=True,timeout=15)
    source=service.create('alice','project','input','Video',blob=video.read_bytes(),mime='video/mp4')
    context=input_context(service,'alice',source['id'])
    assert 1<=len(context['attachments'])<=4
    assert all(item['mime']=='image/jpeg' for item in context['attachments'])


def test_pdf_expansion_is_bounded_without_changing_other_parsers(tmp_path):
    from pypdf import PdfWriter,get_configuration
    from pypdf.generic import NameObject,DecodedStreamObject,DictionaryObject
    service=ArtifactService(tmp_path/'artifacts');writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
    contents=DecodedStreamObject();contents.set_data(b' ' * 8_000_001)
    page[NameObject('/Contents')]=writer._add_object(contents.flate_encode())
    blob=io.BytesIO();writer.write(blob)
    original=get_configuration()
    item=service.create('alice','project','input','Expanded PDF',blob=blob.getvalue(),mime='application/pdf')
    with pytest.raises(HTTPException,match='PDF expanded contents'):input_context(service,'alice',item['id'])
    assert get_configuration()==original
